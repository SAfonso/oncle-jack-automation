import datetime
from pathlib import Path

import openpyxl

from oncle_jack.foto_url import validate_foto_url
from oncle_jack.models import CalendarEntry, Comico, Event

_MESES = [
    "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre",
]

_COL_ESTADO = 14  # columna "Estado" en la hoja Eventos (1-indexada)


def parse_fecha(valor, fila: int) -> datetime.date:
    """Convierte la celda Fecha (datetime/date de Excel o texto dd/mm/yyyy) a date."""
    if isinstance(valor, datetime.datetime):
        return valor.date()
    if isinstance(valor, datetime.date):
        return valor
    if isinstance(valor, str):
        try:
            return datetime.datetime.strptime(valor.strip(), "%d/%m/%Y").date()
        except ValueError:
            pass
    raise ValueError(
        f"fecha invalida en fila {fila} de Eventos: {valor!r} (se espera fecha de Excel o texto dd/mm/yyyy)"
    )


def _mes_de_fecha(fecha: datetime.date) -> str:
    return _MESES[fecha.month - 1]


def load_workbook(path: str | Path, mes: str | None = None) -> tuple[list[CalendarEntry], list[Event]]:
    """Carga calendario y eventos. Si se indica `mes`, valida las fotos de los
    eventos de ese mes (SPEC 5.5); el resto de filas no se validan."""
    wb = openpyxl.load_workbook(path, data_only=True)

    calendario = []
    for mes_cal, sabor, tipo in wb["Calendario_Sabores"].iter_rows(min_row=2, values_only=True):
        if mes_cal is None:
            continue
        calendario.append(CalendarEntry(mes=mes_cal, sabor=sabor, fijo=bool(tipo and "Fijo" in tipo)))

    eventos = []
    for n_fila, row in enumerate(wb["Eventos"].iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None for v in row):
            continue
        fecha = parse_fecha(row[0], n_fila)
        (
            _fecha, lugar, mc_1, mc_2,
            c1_nombre, c1_foto, c2_nombre, c2_foto,
            c3_nombre, c3_foto, c4_nombre, c4_foto,
            sabor_override, estado, _notas,
        ) = row[:15]
        if mes is not None and _mes_de_fecha(fecha) == mes:
            fotos = (c1_foto, c2_foto, c3_foto, c4_foto)
            c1_foto, c2_foto, c3_foto, c4_foto = (
                validate_foto_url(f, fecha.strftime("%d/%m/%Y"), n, f"Comico_{n}_Foto_URL")
                for n, f in enumerate(fotos, 1)
            )
        comicos = [
            Comico(nombre=c1_nombre, foto_url=c1_foto),
            Comico(nombre=c2_nombre, foto_url=c2_foto),
            Comico(nombre=c3_nombre, foto_url=c3_foto),
            Comico(nombre=c4_nombre, foto_url=c4_foto),
        ]
        eventos.append(
            Event(
                fecha=fecha,
                mes=_mes_de_fecha(fecha),
                lugar=lugar,
                mc_1=mc_1,
                mc_2=mc_2,
                comicos=comicos,
                sabor_override=sabor_override,
                estado=estado,
            )
        )

    return calendario, eventos


def save_calendar(path: str | Path, calendario: list[CalendarEntry]) -> None:
    wb = openpyxl.load_workbook(path)
    ws = wb["Calendario_Sabores"]
    por_mes = {entry.mes: entry for entry in calendario}
    for row in ws.iter_rows(min_row=2):
        mes_cell, sabor_cell = row[0], row[1]
        entry = por_mes.get(mes_cell.value)
        if entry is not None:
            sabor_cell.value = entry.sabor
    wb.save(path)


def mark_generated(path: str | Path, evento: Event) -> None:
    wb = openpyxl.load_workbook(path)
    ws = wb["Eventos"]
    for row in ws.iter_rows(min_row=2):
        try:
            fila_fecha = parse_fecha(row[0].value, row[0].row)
        except ValueError:
            continue
        if fila_fecha == evento.fecha:
            ws.cell(row=row[0].row, column=_COL_ESTADO, value="Generado")
            break
    wb.save(path)
