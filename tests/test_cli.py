import json
import shutil
from pathlib import Path

import pytest
from click.testing import CliRunner

from oncle_jack.cli import cli
from oncle_jack.excel_source import load_workbook

FIXTURE = Path(__file__).parent / "fixtures" / "datos_ejemplo.xlsx"


@pytest.fixture
def excel_path(tmp_path):
    destino = tmp_path / "datos.xlsx"
    shutil.copy(FIXTURE, destino)
    return destino


@pytest.fixture
def runner():
    return CliRunner()


def test_resolve_sin_override_usa_el_calendario(runner, excel_path):
    result = runner.invoke(
        cli, ["resolve", "--month", "Noviembre", "--excel", str(excel_path)]
    )

    assert result.exit_code == 0, result.output
    data = json.loads(result.output)
    assert data["mes"] == "Noviembre"
    assert data["sabor"] == "RIE"
    assert data["carpeta_plantilla"] == "Plantillas/RIE"
    assert data["evento"]["fecha"] == "14/11/2026"
    assert data["evento"]["lugar"] == "Carrer de les Roselles, 32"
    assert len(data["evento"]["comicos"]) == 4
    assert data["evento"]["comicos"][0]["nombre"] == "Nombre Cómico 1"


def test_resolve_con_override_intercambia_y_persiste(runner, excel_path):
    result = runner.invoke(
        cli,
        [
            "resolve",
            "--month",
            "Noviembre",
            "--override",
            "Blue",
            "--excel",
            str(excel_path),
        ],
    )

    assert result.exit_code == 0, result.output
    data = json.loads(result.output)
    assert data["sabor"] == "Blue"
    assert data["carpeta_plantilla"] == "Plantillas/Blue"

    calendario, _ = load_workbook(excel_path)
    por_mes = {c.mes: c.sabor for c in calendario}
    assert por_mes["Noviembre"] == "Blue"
    assert por_mes["Abril"] == "RIE"


def test_resolve_con_override_a_mes_fijo_falla_sin_mutar_archivo(runner, excel_path):
    result = runner.invoke(
        cli,
        [
            "resolve",
            "--month",
            "Noviembre",
            "--override",
            "BlackBerry",
            "--excel",
            str(excel_path),
        ],
    )

    assert result.exit_code != 0
    assert "Marzo" in result.output

    calendario, _ = load_workbook(excel_path)
    por_mes = {c.mes: c.sabor for c in calendario}
    assert por_mes["Noviembre"] == "RIE"
    assert por_mes["Marzo"] == "BlackBerry"


def test_reveal_no_necesita_excel(runner):
    result = runner.invoke(cli, ["reveal", "--step", "2"])

    assert result.exit_code == 0, result.output
    data = json.loads(result.output)
    assert data == {"step": 2, "visible_positions": [1, 3], "hidden_positions": [2, 4]}


def test_reveal_paso_invalido_falla(runner):
    result = runner.invoke(cli, ["reveal", "--step", "9"])

    assert result.exit_code != 0


def test_mark_generated_marca_el_evento(runner, excel_path):
    result = runner.invoke(
        cli, ["mark-generated", "--month", "Noviembre", "--excel", str(excel_path)]
    )

    assert result.exit_code == 0, result.output

    _, eventos = load_workbook(excel_path)
    assert eventos[0].estado == "Generado"


def test_resolve_acepta_el_mes_en_minusculas(runner, excel_path):
    result = runner.invoke(
        cli, ["resolve", "--month", "noviembre", "--excel", str(excel_path)]
    )

    assert result.exit_code == 0, result.output
    data = json.loads(result.output)
    assert data["mes"] == "Noviembre"
    assert data["sabor"] == "RIE"


def test_resolve_mes_sin_evento_falla(runner, excel_path):
    result = runner.invoke(
        cli, ["resolve", "--month", "Enero", "--excel", str(excel_path)]
    )

    assert result.exit_code != 0


def test_resolve_foto_invalida_drive_url_falla(runner, excel_path):
    """Modifica fixture: cambia foto del cómico 1 a una URL de Drive (inválida).
    resolve debe salir con exit != 0 y mensaje con fecha, posición, columna y valor."""
    import openpyxl

    # Modificar el fixture temporal: reemplazar foto del cómico 1 por Drive URL
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Eventos"]
    ws["F2"].value = "https://drive.google.com/uc?id=invalid"  # Comico_1_Foto_URL (col F = 6)
    wb.save(excel_path)

    result = runner.invoke(
        cli, ["resolve", "--month", "Noviembre", "--excel", str(excel_path)]
    )

    assert result.exit_code != 0
    # Verificar que el error contiene fecha, posición, columna y valor inválido
    assert "14/11/2026" in result.output  # fecha del evento
    assert "Comico_1_Foto_URL" in result.output  # columna
    assert "drive.google.com" in result.output  # valor inválido


def test_resolve_foto_vacia_falla(runner, excel_path):
    """Modifica fixture: deja vacía la foto del cómico 2.
    resolve debe salir con exit != 0 y mensaje con fecha, posición, columna y valor."""
    import openpyxl

    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Eventos"]
    ws["H2"].value = None  # Comico_2_Foto_URL (col H = 8)
    wb.save(excel_path)

    result = runner.invoke(
        cli, ["resolve", "--month", "Noviembre", "--excel", str(excel_path)]
    )

    assert result.exit_code != 0
    assert "14/11/2026" in result.output  # fecha del evento
    assert "Comico_2_Foto_URL" in result.output  # columna
    assert "2" in result.output  # posición (cómico 2)


def test_resolve_foto_valida_contiene_canva_link_en_json(runner, excel_path):
    """Con fixture válido (sin modificar), resolve sale con exit 0 y JSON
    contiene foto_url con canva.link."""
    result = runner.invoke(
        cli, ["resolve", "--month", "Noviembre", "--excel", str(excel_path)]
    )

    assert result.exit_code == 0, result.output
    data = json.loads(result.output)
    # Verificar que al menos uno de los cómicos tiene foto_url con canva.link
    fotos = [c["foto_url"] for c in data["evento"]["comicos"]]
    assert any("canva.link" in f or "canva.com" in f for f in fotos), f"Ninguna foto contiene URL canva: {fotos}"


def test_mark_generated_funciona_con_foto_invalida(runner, excel_path):
    """mark-generated NO debe validar fotos. Incluso con foto inválida,
    debe marcar el evento como Generado sin error."""
    import openpyxl

    # Modificar el fixture temporal: reemplazar foto del cómico 1 por Drive URL
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Eventos"]
    ws["F2"].value = "https://drive.google.com/uc?id=invalid"  # Comico_1_Foto_URL (col F = 6)
    wb.save(excel_path)

    result = runner.invoke(
        cli, ["mark-generated", "--month", "Noviembre", "--excel", str(excel_path)]
    )

    # mark-generated NO debe fallar por foto inválida
    assert result.exit_code == 0, result.output

    # Verificar que el evento fue marcado como Generado
    _, eventos = load_workbook(excel_path)
    assert eventos[0].estado == "Generado"


def test_mark_generated_json_fecha_dd_mm_yyyy_con_fila_datetime(runner, excel_path):
    result = runner.invoke(
        cli, ["mark-generated", "--month", "Diciembre", "--excel", str(excel_path)]
    )
    assert result.exit_code == 0, result.output
    assert json.loads(result.output)["fecha"] == "12/12/2026"
    _, eventos = load_workbook(excel_path)
    assert eventos[1].estado == "Generado"


def test_resolve_mes_valido_ok_aunque_otros_meses_tengan_campos_vacios(runner, excel_path):
    # El fixture ya trae Diciembre con nombres/fotos vacios.
    result = runner.invoke(
        cli, ["resolve", "--month", "Noviembre", "--excel", str(excel_path)]
    )
    assert result.exit_code == 0, result.output


def test_resolve_mes_con_nombre_vacio_falla_con_fecha_posicion_columna(runner, excel_path):
    result = runner.invoke(
        cli, ["resolve", "--month", "Diciembre", "--excel", str(excel_path)]
    )
    assert result.exit_code != 0
    assert result.exception is None or isinstance(result.exception, SystemExit)
    assert "12/12/2026" in result.output
    assert "comico 1" in result.output
    assert "Comico_1_Nombre" in result.output
    assert "Traceback" not in result.output


def test_resolve_override_no_persiste_si_evento_invalido(runner, excel_path):
    antes, _ = load_workbook(excel_path)
    result = runner.invoke(
        cli,
        ["resolve", "--month", "Diciembre", "--override", "Blue", "--excel", str(excel_path)],
    )
    assert result.exit_code != 0
    despues, _ = load_workbook(excel_path)
    assert antes == despues


def test_mark_generated_no_valida_nombres_vacios(runner, excel_path):
    result = runner.invoke(
        cli, ["mark-generated", "--month", "Diciembre", "--excel", str(excel_path)]
    )
    assert result.exit_code == 0, result.output
