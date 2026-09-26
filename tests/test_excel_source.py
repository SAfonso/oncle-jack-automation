import shutil
from pathlib import Path

import openpyxl
import pytest

from oncle_jack.excel_source import load_workbook, mark_generated, save_calendar

FIXTURE = Path(__file__).parent / "fixtures" / "datos_ejemplo.xlsx"


@pytest.fixture
def excel_path(tmp_path):
    destino = tmp_path / "datos.xlsx"
    shutil.copy(FIXTURE, destino)
    return destino


def test_load_workbook_lee_calendario_y_eventos(excel_path):
    calendario, eventos = load_workbook(excel_path)

    octubre = next(c for c in calendario if c.mes == "Octubre")
    assert octubre.sabor == "Original"
    assert octubre.fijo is True

    marzo = next(c for c in calendario if c.mes == "Marzo")
    assert marzo.sabor == "BlackBerry"
    assert marzo.fijo is True

    noviembre = next(c for c in calendario if c.mes == "Noviembre")
    assert noviembre.sabor == "RIE"
    assert noviembre.fijo is False

    assert len(eventos) == 1
    evento = eventos[0]
    assert evento.fecha == "14/11/2026"
    assert evento.mes == "Noviembre"
    assert evento.lugar == "Carrer de les Roselles, 32"
    assert evento.mc_1 == "Mercè Copado"
    assert evento.mc_2 == "Sergio Afonso"
    assert evento.estado == "Pendiente"
    assert evento.sabor_override is None
    assert len(evento.comicos) == 4
    assert evento.comicos[0].nombre == "Nombre Cómico 1"
    assert evento.comicos[1].nombre == "Nombre Cómico 2"
    assert evento.comicos[3].foto_url == "https://canva.link/fixfoto4"


def test_save_calendar_escribe_solo_la_hoja_calendario(excel_path):
    calendario, eventos_antes = load_workbook(excel_path)
    for entry in calendario:
        if entry.mes == "Noviembre":
            entry.sabor = "Blue"

    save_calendar(excel_path, calendario)

    calendario_recargado, eventos_despues = load_workbook(excel_path)
    noviembre = next(c for c in calendario_recargado if c.mes == "Noviembre")
    assert noviembre.sabor == "Blue"
    # no ha tocado el resto del archivo
    assert eventos_despues == eventos_antes


def test_mark_generated_marca_el_evento_correspondiente(excel_path):
    _, eventos = load_workbook(excel_path)
    evento = eventos[0]
    assert evento.estado == "Pendiente"

    mark_generated(excel_path, evento)

    _, eventos_recargados = load_workbook(excel_path)
    assert eventos_recargados[0].estado == "Generado"


def test_no_muta_el_fixture_original(excel_path):
    calendario, _ = load_workbook(excel_path)
    for entry in calendario:
        entry.sabor = "Cambiado"
    save_calendar(excel_path, calendario)

    calendario_original, _ = load_workbook(FIXTURE)
    assert all(c.sabor != "Cambiado" for c in calendario_original)


def test_fixture_cabecera_col7_es_comico_2_nombre():
    ws = openpyxl.load_workbook(FIXTURE)["Eventos"]
    assert ws.cell(row=1, column=7).value == "Comico_2_Nombre"


def test_fixture_sin_rastro_de_drive():
    wb = openpyxl.load_workbook(FIXTURE)
    for ws in wb:
        for row in ws.iter_rows(values_only=True):
            assert not any("drive.google" in str(v) for v in row if v)


def test_load_workbook_con_mes_valida_evento_valido(excel_path):
    _, eventos = load_workbook(excel_path, mes="Noviembre")
    assert eventos[0].comicos[0].foto_url == "https://canva.link/fixfoto1"


@pytest.mark.parametrize(
    "posicion,columna,col_idx,valor",
    [
        (1, "Comico_1_Foto_URL", 6, None),
        (2, "Comico_2_Foto_URL", 8, "https://drive.google.com/uc?id=X"),
        (3, "Comico_3_Foto_URL", 10, "http://canva.link/abc"),
        (4, "Comico_4_Foto_URL", 12, "https://ejemplo.com/foto.png"),
    ],
)
def test_load_workbook_con_mes_rechaza_foto_invalida(excel_path, posicion, columna, col_idx, valor):
    wb = openpyxl.load_workbook(excel_path)
    wb["Eventos"].cell(row=2, column=col_idx).value = valor
    wb.save(excel_path)

    with pytest.raises(ValueError) as exc:
        load_workbook(excel_path, mes="Noviembre")
    msg = str(exc.value)
    assert "14/11/2026" in msg
    assert f"comico {posicion}" in msg
    assert columna in msg
    assert repr(valor) in msg


def test_load_workbook_sin_mes_no_valida_filas_futuras(excel_path):
    wb = openpyxl.load_workbook(excel_path)
    wb["Eventos"].cell(row=2, column=6).value = None
    wb.save(excel_path)

    _, eventos = load_workbook(excel_path)
    assert eventos[0].comicos[0].foto_url is None


def test_load_workbook_con_mes_ignora_fotos_vacias_de_otros_meses(excel_path):
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Eventos"]
    for c in range(1, 16):
        ws.cell(row=3, column=c).value = ws.cell(row=2, column=c).value
    ws.cell(row=3, column=1).value = "12/12/2026"
    ws.cell(row=3, column=6).value = None
    wb.save(excel_path)

    _, eventos = load_workbook(excel_path, mes="Noviembre")
    assert len(eventos) == 2
